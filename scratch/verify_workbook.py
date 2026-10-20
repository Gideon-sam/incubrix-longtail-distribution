import openpyxl

wb = openpyxl.load_workbook('Assessment4_Gideon_Samuel_Final.xlsx', data_only=True)
for sh in wb.sheetnames:
    ws = wb[sh]
    print(f'Sheet: {sh} ({ws.max_row}r x {ws.max_column}c)')

ws2 = wb['2_Must_haves']
print('\nMust-haves:')
for r in range(3, 14):
    a = ws2.cell(r, 1).value
    b = ws2.cell(r, 2).value
    c = ws2.cell(r, 3).value
    if a:
        print(f'  MH{r-2}: {str(a)[:40]} | {b} | {str(c)[:50] if c else ""}')

ws3 = wb['3_Levels']
print('\nLevels:')
for r in range(3, 22):
    a = ws3.cell(r, 1).value
    b = ws3.cell(r, 2).value
    c = ws3.cell(r, 3).value
    if a:
        print(f'  {str(a)[:40]} | {b} | {str(c)[:40] if c else ""}')

ws4 = wb['4_Integration']
print('\nIntegration:')
for r in range(3, 16):
    a = ws4.cell(r, 1).value
    b = ws4.cell(r, 2).value
    c = ws4.cell(r, 3).value
    if a:
        print(f'  {str(a)[:40]} | {b} | {str(c)[:40] if c else ""}')
